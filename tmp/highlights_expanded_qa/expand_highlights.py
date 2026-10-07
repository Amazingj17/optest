from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(r'D:\programing\python\optest')
SOURCE=BASE/'docs/项目说明书_修订版.docx'
OUTPUT=BASE/'docs/项目说明书_亮点详解版.docx'
ASSETS=BASE/'presentation/assets'
doc=Document(SOURCE)
start=next(p for p in doc.paragraphs if p.text=='1.3 项目亮点')
end=next(p for p in doc.paragraphs if p.text=='第 2 章 技术理论')
body_style=next(p for p in doc.paragraphs if p.text.startswith('项目据此确定了三项工程需求'))
heading_style=next(p for p in doc.paragraphs if p.text=='2.1.1 问题建模与统一调度语义')
original_ppr=deepcopy(body_style._p.pPr)
heading_ppr=deepcopy(heading_style._p.pPr)
original_end=deepcopy(end._p)
original_tables=[deepcopy(t._tbl) for t in doc.tables]

current=start._p.getnext()
while current is not end._p:
    nxt=current.getnext()
    current.getparent().remove(current)
    current=nxt

added=[]
def insert(p):
    end._p.addprevious(p._p)
    added.append(p)
    return p
def paragraph(value,heading=False):
    p=doc.add_paragraph()
    p._p.insert(0,deepcopy(heading_ppr if heading else original_ppr))
    p.add_run(value)
    if heading:p.paragraph_format.keep_with_next=True
    return insert(p)

def figure(path,caption,crop=None):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    ind=p._p.get_or_add_pPr().get_or_add_ind();ind.set(qn('w:firstLineChars'),'0')
    p.paragraph_format.space_before=Pt(6)
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1.0
    picture=p.add_run().add_picture(str(path),width=Inches(6.53))
    if crop:
        # Native Word picture cropping keeps the original asset bytes intact.
        top,bottom=crop
        blipfill=picture._inline.xpath('.//pic:blipFill')[0]
        src=OxmlElement('a:srcRect');src.set('t',str(top));src.set('b',str(bottom))
        blipfill.insert(1,src)
        picture.height=int(picture.height*(1-(top+bottom)/100000))
    picture._inline.docPr.set('descr',caption)
    insert(p)
    c=doc.add_paragraph()
    c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.first_line_indent=Pt(0)
    ci=c._p.get_or_add_pPr().get_or_add_ind();ci.set(qn('w:firstLineChars'),'0')
    c.paragraph_format.space_before=Pt(3)
    c.paragraph_format.space_after=Pt(6)
    c.paragraph_format.line_spacing=1.2
    c.paragraph_format.keep_together=True
    r=c.add_run(caption);r.font.size=Pt(10.5)
    insert(c)

paragraph('项目亮点贯穿任务与资源建模、分层学习、候选搜索和结果交付。面对任务依赖、节点竞争与通信等待共同作用的调度问题，系统先形成合法决策，再通过完整时间线判断方案质量，并保留配置、模型和候选来源，使结果能够解释和复核。以下从具体机制与实现价值展开说明。')

paragraph('1.3.1 分层模型与启发式残差学习',True)
paragraph('任务排序和资源分配相互影响：先选择哪一个就绪任务，会改变可供节点选择的执行与通信条件；任务放到哪个节点，又会改变后续任务的数据到达和资源等待。项目将决策拆成高层选任务、低层选节点，使两层分别处理自己的动作集合，同时通过同一调度环境交换状态。')
paragraph('高层 LSTM 编码当前任务特征序列，为每个任务生成动作分数；低层 GAT 在选中任务的条件下编码资源图，综合节点速度、带宽与时延、预计开始及完成时间生成节点分数。ready_mask 限制任务的前驱就绪条件，node_mask 检查节点设备与容量是否可行。选中动作后，插入式模拟器在满足数据到达的空闲区间放置任务，并返回新观测和奖励。这里的 LSTM 每次重新编码当前序列，不跨调度决策保留历史隐状态。')
figure(ASSETS/'hierarchical-scheduling-academic.png','图1-2 分层决策与启发式残差的协同流程',crop=(25000,25000))
paragraph('如图1-2所示，模型在 HEFT rank 和 EFT 先验上学习有界修正。正式配置使用权重 8 的先验与幅度上限 0.5 的 tanh 残差，actor 末层零初始化，使初始决策继承启发式偏好。训练先固定低层 EFT 规则、更新高层，再进入两层联合训练，并在验证选模中保留初始与最佳检查点。这套设计将初始调度能力、学习修正与阶段控制分开记录，便于定位更新耦合或后期退化；具体网络与参数见第 2 章。')

paragraph('1.3.2 实际关键链与任务块搜索',True)
paragraph('DAG 中依赖路径最长，不一定意味着实际调度中等待最多。任务还可能被同一节点上已经排定的任务阻塞，跨节点传输也会推迟数据就绪。项目从已经生成的调度时间线回溯关键链，同时考虑前驱数据到达和节点前序任务占用，定位真正限制整体完成时刻的任务。')
paragraph('搜索围绕实际关键链尝试节点迁移与次序调整，并对短关键块进行共同迁移。单个任务移到快节点时，新增通信可能抵消计算收益；将相互依赖的任务块一起迁移，则可能避免块内跨节点传输。估计代价用于筛选移动，最终收益由完整候选重放后的真实 makespan 判断，不能仅凭局部估计宣布调度改进。')
figure(ASSETS/'imagegen_revision/search.png','图1-3 从实际关键链定位到任务块迁移与完整候选评分')
paragraph('一个最小机制案例可以说明共同迁移的作用：任务 a、b 依次执行，工作量分别为 4、2，慢节点速度为 1，快节点速度为 2，跨节点传输耗时为 10。两任务都在慢节点时 makespan 为 6；只迁移 a 或 b 时分别变为 14、15；共同迁到快节点后变为 3。对应测试验证了单任务迁移均变差，而关键块移动能够改善的情形。该案例说明搜索为何需要联合移动，不代表全验证集的平均收益；实现与测试见第 5.3 节。')

paragraph('1.3.3 动态状态隔离与前缀复用',True)
paragraph('搜索需要比较多个完整方案，若每次只改动后半段也从空状态重放全部任务，未变化的前缀会被反复计算。直接让多个候选共享可变时间线又会造成分支污染：分支 A 的放置改变父状态后，分支 B 的等待时间和评分也会随之变化，候选比较便失去独立性。')
paragraph('项目将静态场景与动态调度状态分开管理。任务 DAG、资源矩阵等静态信息可以共享，节点时间线、任务放置记录、就绪集合和决策顺序则在克隆时复制。搜索按父候选与前缀深度保存快照，复用改动位置之前的调度，再克隆快照并重新模拟后缀。切换父候选时重建前缀缓存，避免把不同方案中同一深度的状态误当作相同前缀。')
figure(ASSETS/'imagegen_revision/reuse.png','图1-4 共享静态信息与共同前缀 克隆各分支动态状态')
paragraph('例如，两个候选均保留前四次决策，只在后续改变任务节点或次序，就可以复用这四次决策对应的前缀快照；两条后缀仍各自在独立状态上执行。测试检查分支放置不改变父状态，缓存开关不改变候选结果，并用模拟调用计数验证重复重放是否减少。亮点在于复用计算时仍保持精确时间线与候选隔离；实际墙钟收益需要在相同场景与搜索预算下另行测量。')

paragraph('1.3.4 完整候选校验与独立组合',True)
paragraph('系统将 HEFT、独立搜索和冻结 HRL 作为三路完整候选。三路在相同场景语义下分别建立环境，按顺序生成完整调度，前一路结果不作为后一路搜索起点；独立搜索的 learned_policy 为 None，HRL 不参与 beam 分支展开。图1-5的分支表示候选与校验之间的数据关系，不表示当前实现并发执行。')
paragraph('比较前先检查任务覆盖、依赖和资源约束、时间线以及有限正 makespan，并确认各分支采用一致的 HEFT 参考分母。全部候选通过后取 makespan 最小者，精确持平按 HEFT、搜索、HRL 的顺序选择。选中的决策序列在最终环境逐动作重放，每一步重新检查 ready_mask 与 node_mask；候选失败或参考分母不一致时明确报错并清空旧计划，避免残留结果被当作本次成功输出。')
figure(ASSETS/'system-architecture-academic.png','图1-5 独立完整候选的统一校验 择优与结果复核',crop=(5000,5000))
paragraph('这种组合把“提出方案”与“验证方案”连接起来：学习模块和搜索模块可以独立评价，最终输出则有候选分数、选择来源和重放记录可查。在候选均成功、合法且采用同一确定性目标时，选中方案不会差于其中任何一个已完成候选。这一关系提供已生成方案之间的质量保护，仍需计入全部候选生成、校验、选择与重放成本；核心过程见算法2-1。')

paragraph('1.3.5 可复现评测与双开源系统运行',True)
paragraph('项目通过统一 Scenario、执行通信模型、动作约束和插入时间线，支持 HEFT、学习策略、独立搜索与组合方法在同一输入上比较。数据按基础 DAG 隔离划分，正式批次使用三个随机种子，保存训练预算、检查点、配置、逐场景结果及来源记录。演示页面提供策略切换、任务依赖查看、时间线回放和 JSON 导出，使评测结果能够从汇总指标追溯到具体调度。')
paragraph('固定 108 个验证场景、54 个基础 DAG 上，组合方法的 mean_ratio 为 0.918223 ± 0.000498，“±”为种子间样本标准差；相对 HEFT，平均归一化 makespan 降低约 8.18%，归档调度合法率为 100%。相同批次中，独立搜索 mean_ratio 为 0.918239，纯 HRL 为 0.999058，说明主要收益来自搜索；组合平均 CPU 推理耗时约 5.297 秒/场景，质量提升伴随额外求解成本。各方法的表现与学习增量在第 4 章分别报告。')
paragraph('运行交付方面，openEuler 与 openKylin 的 x86_64 CPU 虚拟机均保存了 2 回合短训练、400 次任务决策和 108 场景验证记录，HRL 与 HEFT 合法率均为 100%。openEuler 的 188 项测试和六阶段验收通过，重载评估与训练后结果一致；openKylin 已归档数值结果及测试、重载完成说明，过程日志仍需补齐。双系统记录提供对应环境下的可运行证据，与三种子正式性能实验分别陈述；配置和验收细节见第 4、5 章。')

doc.save(OUTPUT)
# Preserve native math and all pre-existing tables outside the replaced section.
NS={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(SOURCE) as a,ZipFile(OUTPUT) as b:
    r1=E.fromstring(a.read('word/document.xml'));r2=E.fromstring(b.read('word/document.xml'))
    math1=[''.join(m.xpath('.//m:t/text()',namespaces=NS))for m in r1.xpath('.//m:oMath',namespaces=NS)]
    math2=[''.join(m.xpath('.//m:t/text()',namespaces=NS))for m in r2.xpath('.//m:oMath',namespaces=NS)]
    assert math1==math2 and len(math2)==7
    assert len(r1.xpath('.//w:tbl',namespaces=NS))==len(r2.xpath('.//w:tbl',namespaces=NS))
    for n in a.namelist():
        if n.startswith('word/media/'):assert a.read(n)==b.read(n)
print('Expanded 1.3 into five subsections with four diagrams. Native formulas and existing media retained.')
print(OUTPUT)
